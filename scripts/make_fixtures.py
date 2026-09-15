"""Generates small real-world test documents for the Rust import tests.

Usage:  python scripts/make_fixtures.py <output-dir>
Then:   SIFTQDA_FIXTURES=<output-dir> cargo test   (from src-tauri)
Requires python-docx, openpyxl and fpdf (or fpdf2).
"""

import sys
from pathlib import Path

from docx import Document
from fpdf import FPDF
from openpyxl import Workbook


def make_docx(path: Path) -> None:
    doc = Document()
    doc.core_properties.author = "Test Researcher"
    doc.core_properties.title = "Participant 7"
    doc.add_heading("Interview with Participant 7", level=1)
    doc.add_paragraph("Interviewer: How has the weather affected you?")
    doc.add_paragraph("Participant: Honestly, the heat waves make me anxious every summer.")
    doc.add_paragraph("Sleep problems", style="List Bullet")
    doc.add_paragraph("Worry about my children", style="List Bullet")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Age"
    table.cell(0, 1).text = "34"
    table.cell(1, 0).text = "Region"
    table.cell(1, 1).text = "Coastal"
    doc.save(path)


def make_xlsx(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Survey"
    ws.append(["ID", "Age", "Region", "Open answer"])
    ws.append(["P1", 34, "Coastal", "I worry about floods."])
    ws.append(["P2", 51, "Inland", "Summers feel longer."])
    ws.append([None, None, None, None])
    ws.append(["P3", 27, "Coastal", "Heat keeps me awake."])
    wb.save(path)


def make_pdf(path: Path) -> None:
    pdf = FPDF()
    pdf.set_font("Helvetica", size=12)
    pdf.add_page()
    pdf.multi_cell(0, 8, "Page one talks about rising temperatures\nand coastal flooding.")
    pdf.add_page()
    pdf.multi_cell(0, 8, "Page two covers community responses.")
    pdf.output(str(path))


def make_odt(path: Path) -> None:
    """A minimal but valid OpenDocument text file, built by hand (no odfpy needed)."""
    import zipfile

    ns = (
        'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
        'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" '
        'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
        'xmlns:dc="http://purl.org/dc/elements/1.1/" '
        'xmlns:meta="urn:oasis:names:tc:opendocument:xmlns:meta:1.0"'
    )
    content = f"""<?xml version="1.0" encoding="UTF-8"?>
<office:document-content {ns} office:version="1.3"><office:body><office:text>
<text:h text:outline-level="1">Focus group 2</text:h>
<text:p>Moderator: What worries you<text:s text:c="2"/>most?</text:p>
<text:p>Speaker A: <text:span>Losing our crops</text:span> to drought.<text:note><text:note-body><text:p>footnote</text:p></text:note-body></text:note></text:p>
<text:list><text:list-item><text:p>Water shortages</text:p></text:list-item></text:list>
<table:table><table:table-row><table:table-cell><text:p>Village</text:p></table:table-cell><table:table-cell><text:p>North</text:p></table:table-cell></table:table-row></table:table>
</office:text></office:body></office:document-content>"""
    meta = f"""<?xml version="1.0" encoding="UTF-8"?>
<office:document-meta {ns}><office:meta><meta:initial-creator>Test Researcher</meta:initial-creator></office:meta></office:document-meta>"""
    manifest = """<?xml version="1.0" encoding="UTF-8"?>
<manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0" manifest:version="1.3">
<manifest:file-entry manifest:full-path="/" manifest:media-type="application/vnd.oasis.opendocument.text"/>
<manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/>
<manifest:file-entry manifest:full-path="meta.xml" manifest:media-type="text/xml"/>
</manifest:manifest>"""
    with zipfile.ZipFile(path, "w") as z:
        z.writestr(zipfile.ZipInfo("mimetype"), "application/vnd.oasis.opendocument.text")
        z.writestr("META-INF/manifest.xml", manifest, zipfile.ZIP_DEFLATED)
        z.writestr("content.xml", content, zipfile.ZIP_DEFLATED)
        z.writestr("meta.xml", meta, zipfile.ZIP_DEFLATED)


def main() -> None:
    out = Path(sys.argv[1] if len(sys.argv) > 1 else "fixtures")
    out.mkdir(parents=True, exist_ok=True)
    make_docx(out / "interview.docx")
    make_xlsx(out / "survey.xlsx")
    make_pdf(out / "report.pdf")
    make_odt(out / "focus-group.odt")
    (out / "notes.txt").write_text("﻿First line\r\n\r\nSecond line with ünïcode\n", encoding="utf-8")
    print(f"fixtures written to {out.resolve()}")


if __name__ == "__main__":
    main()
